"""Customer delivery-schedule import with preview, validation and apply (BRD FR-SAL-03).

A customer's schedule arrives as a spreadsheet (.xlsx), a CSV, or text pasted from an e-mail.
Uploading parses it and shows every row matched to the customer's order (by the customer's PO
number) and our item (by the customer's part number), with what applying would do: CREATE a new
delivery date or UPDATE an existing one, or an error saying why the row cannot be used. Nothing
changes until the import is applied, and an applied import cannot be applied twice. Rows are
checked again when applying, because orders may have changed since the preview. The scheduled
quantity of a line can never exceed what was ordered. Needs sales.order.
"""

import base64
import binascii
import csv
import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, status
from openpyxl import load_workbook
from pydantic import BaseModel, Field, model_validator

from heyyerp_api.masters.common import (
    Ctx,
    Page,
    Paging,
    Strict,
    audit,
    count,
    not_found,
    one,
    requires,
    requires_module,
    rows,
)

router = APIRouter(prefix="/sales", tags=["sales"])
manage = requires("sales.order")
read = requires_module("sales")
MAX_ROWS = 2000
MAX_BYTES = 5_000_000
ALIASES = {
    "customer_po_no": ("customer_po_no", "po", "po_no", "po number", "po no", "customer po",
                       "customer po no", "order no"),
    "customer_part_no": ("customer_part_no", "part_no", "part no", "part number", "part",
                         "customer part no", "customer part", "item"),
    "due_date": ("due_date", "due date", "date", "delivery date", "delivery", "due"),
    "quantity": ("quantity", "qty", "scheduled_qty", "schedule qty", "required qty"),
}
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y")


class ImportIn(Strict):
    customer_id: UUID
    source_type: Literal["XLSX", "CSV", "EMAIL", "MANUAL"]
    text: str | None = Field(default=None, max_length=1_000_000)
    content_base64: str | None = Field(default=None, max_length=7_000_000)

    @model_validator(mode="after")
    def _one_source(self) -> "ImportIn":
        if (self.text is None) == (self.content_base64 is None):
            raise ValueError("Send either text or content_base64")
        if self.source_type == "XLSX" and self.content_base64 is None:
            raise ValueError("An .xlsx import needs content_base64")
        return self


class ApplyIn(Strict):
    skip_errors: bool = False


class PreviewRow(BaseModel):
    row_no: int
    customer_po_no: str | None
    customer_part_no: str | None
    due_date: date | None
    quantity: Decimal | None
    so_no: str | None = None
    so_line_id: UUID | None = None
    item_code: str | None = None
    action: str  # CREATE, UPDATE or ERROR
    error: str | None = None


class ImportSummary(BaseModel):
    import_id: UUID
    customer_id: UUID | None
    customer_code: str | None
    source_type: str
    status: str
    created_at: datetime
    applied_at: datetime | None
    rows_total: int
    rows_ok: int
    rows_error: int


class ImportOut(ImportSummary):
    rows: list[PreviewRow]


# ------------------------------------------------------------------ parsing
def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_quantity(value: Any) -> Decimal | None:
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None
    return number if number > 0 and number < Decimal(1_000_000_000) else None


def to_table(body: ImportIn) -> list[list[Any]]:
    """Header row plus data rows, from an xlsx file or delimited text."""
    if body.source_type == "XLSX":
        try:
            raw = base64.b64decode(body.content_base64 or "", validate=True)
        except (binascii.Error, ValueError) as exc:
            raise HTTPException(422, "The file content is not valid base64") from exc
        if len(raw) > MAX_BYTES:
            raise HTTPException(422, "The file is larger than 5 MB")
        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
            return [list(r) for r in sheet.iter_rows(values_only=True)]
        except Exception as exc:  # openpyxl raises many kinds of errors for a bad file
            raise HTTPException(422, "The file is not a readable .xlsx workbook") from exc
    if body.text is not None:
        text = body.text
    else:
        try:
            text = base64.b64decode(body.content_base64 or "", validate=True).decode("utf-8-sig")
        except (binascii.Error, ValueError) as exc:
            raise HTTPException(422, "The file is not valid UTF-8 text") from exc
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        raise HTTPException(422, "The import has no rows")
    delimiter = max((",", "\t", ";", "|"), key=lines[0].count)
    return [list(r) for r in csv.reader(lines, delimiter=delimiter)]


def parse(body: ImportIn) -> list[dict[str, Any]]:
    table = [r for r in to_table(body) if any(str(c or "").strip() for c in r)]
    if len(table) < 2:
        raise HTTPException(422, "The import needs a header row and at least one data row")
    if len(table) - 1 > MAX_ROWS:
        raise HTTPException(422, f"An import may have at most {MAX_ROWS} rows")
    header = [str(c or "").strip().lower() for c in table[0]]
    index: dict[str, int] = {}
    for key, names in ALIASES.items():
        for position, name in enumerate(header):
            if name in names and key not in index:
                index[key] = position
    missing = [k for k in ALIASES if k not in index]
    if missing:
        raise HTTPException(422, f"Columns not found: {', '.join(missing)}")
    parsed = []
    for number, row in enumerate(table[1:], start=2):
        def cell(key: str, row: list[Any] = row) -> Any:
            return row[index[key]] if index[key] < len(row) else None

        po, part = cell("customer_po_no"), cell("customer_part_no")
        parsed.append({
            "row_no": number,
            "customer_po_no": str(po).strip() if po not in (None, "") else None,
            "customer_part_no": str(part).strip() if part not in (None, "") else None,
            "due_date": parse_date(cell("due_date")),
            "quantity": parse_quantity(cell("quantity")),
        })
    return parsed


# ------------------------------------------------------------------ matching
async def evaluate(ctx: Ctx, customer_id: UUID, raw_rows: list[dict[str, Any]]) -> list[PreviewRow]:
    """Match each row to an order line and say what applying it would do."""
    out: list[PreviewRow] = []
    seen: set[tuple[UUID, date]] = set()
    planned: dict[UUID, Decimal] = {}
    for r in raw_rows:
        row = PreviewRow(row_no=r["row_no"], customer_po_no=r["customer_po_no"],
                         customer_part_no=r["customer_part_no"], due_date=r["due_date"],
                         quantity=r["quantity"], action="ERROR")
        out.append(row)
        if not row.customer_po_no or not row.customer_part_no:
            row.error = "The PO number and part number are required"
            continue
        if row.due_date is None:
            row.error = "The delivery date is missing or not a valid date"
            continue
        if row.quantity is None:
            row.error = "The quantity is missing or not a positive number"
            continue
        order = await one(
            ctx, "SELECT so_id, so_no, order_date, status FROM erp.sales_order WHERE customer_id "
            "= :c AND customer_po_no = :p AND status <> 'CANCELLED'",
            {"c": customer_id, "p": row.customer_po_no})
        if order is None:
            row.error = f"No order of this customer has PO number {row.customer_po_no}"
            continue
        row.so_no = order["so_no"]
        if order["status"] not in ("OPEN", "PARTIALLY_SHIPPED"):
            row.error = f"Order {order['so_no']} is {str(order['status']).lower()}"
            continue
        line = await one(
            ctx, "SELECT l.so_line_id, l.ordered_qty, i.code FROM erp.customer_part_xref x JOIN "
            "erp.so_line l ON l.item_id = x.item_id AND l.so_id = :so JOIN erp.item i ON "
            "i.item_id = l.item_id WHERE x.customer_id = :c AND x.customer_part_no = :n",
            {"so": order["so_id"], "c": customer_id, "n": row.customer_part_no})
        if line is None:
            row.error = (f"Part {row.customer_part_no} is not on order {order['so_no']} "
                         "(or has no part-number cross-reference)")
            continue
        row.so_line_id, row.item_code = line["so_line_id"], line["code"]
        if row.due_date < order["order_date"]:
            row.error = "The delivery date is before the order date"
            continue
        if (line["so_line_id"], row.due_date) in seen:
            row.error = "The same part and date appear twice in this file"
            continue
        seen.add((line["so_line_id"], row.due_date))
        existing = await one(
            ctx, "SELECT scheduled_qty, fulfilled_qty FROM erp.so_schedule WHERE so_line_id = :l "
            "AND due_date = :d", {"l": line["so_line_id"], "d": row.due_date})
        if existing is not None and row.quantity < existing["fulfilled_qty"]:
            row.error = f"{existing['fulfilled_qty']} already shipped for that date"
            continue
        others = await one(
            ctx, "SELECT COALESCE(sum(scheduled_qty), 0) AS q FROM erp.so_schedule WHERE "
            "so_line_id = :l AND due_date <> :d", {"l": line["so_line_id"], "d": row.due_date})
        assert others is not None
        total = others["q"] + planned.get(line["so_line_id"], Decimal(0)) + row.quantity
        if total > line["ordered_qty"]:
            row.error = f"The schedule would exceed the ordered quantity {line['ordered_qty']}"
            continue
        planned[line["so_line_id"]] = planned.get(line["so_line_id"], Decimal(0)) + row.quantity
        row.action = "UPDATE" if existing is not None else "CREATE"
    return out


def summarize(import_id: UUID, head: dict[str, Any], preview: list[PreviewRow]) -> ImportOut:
    ok = sum(1 for r in preview if r.action != "ERROR")
    return ImportOut(
        import_id=import_id, customer_id=head["customer_id"], customer_code=head["customer_code"],
        source_type=head["source_type"], status=head["status"], created_at=head["created_at"],
        applied_at=head["applied_at"], rows_total=len(preview), rows_ok=ok,
        rows_error=len(preview) - ok, rows=preview)


IMPORT_SQL = (
    "SELECT s.import_id, s.customer_id, c.code AS customer_code, s.source_type, s.status, "
    "s.created_at, s.applied_at, s.parsed_rows FROM erp.schedule_import s LEFT JOIN erp.customer "
    "c ON c.customer_id = s.customer_id"
)


async def load(ctx: Ctx, import_id: UUID) -> ImportOut | None:
    head = await one(ctx, f"{IMPORT_SQL} WHERE s.import_id = :id", {"id": import_id})
    if head is None:
        return None
    stored = head["parsed_rows"]
    data = json.loads(stored) if isinstance(stored, str) else stored
    return summarize(import_id, head, [PreviewRow(**r) for r in data])


@router.post("/schedule-imports", response_model=ImportOut, status_code=status.HTTP_201_CREATED)
async def upload(body: ImportIn, ctx: Ctx = Depends(manage)) -> ImportOut:
    customer = await one(ctx, "SELECT active FROM erp.customer WHERE customer_id = :id",
                         {"id": body.customer_id})
    if customer is None or not customer["active"]:
        raise HTTPException(422, "The customer does not exist or is inactive")
    raw_rows = parse(body)
    preview = await evaluate(ctx, body.customer_id, raw_rows)
    created = await one(
        ctx, "INSERT INTO erp.schedule_import(tenant_id, uploaded_by, source_type, status, "
        "parsed_rows, customer_id) VALUES (erp.current_tenant_id(), :me, :t, 'PREVIEW', "
        "CAST(:rows AS jsonb), :c) RETURNING import_id",
        {"me": ctx.principal.profile.app_user_id, "t": body.source_type, "c": body.customer_id,
         "rows": json.dumps([r.model_dump(mode="json") for r in preview])})
    assert created is not None
    await audit(ctx, "schedule_import.upload", "schedule_import", created["import_id"],
                {"rows": len(preview), "errors": sum(1 for r in preview if r.action == "ERROR")})
    result = await load(ctx, created["import_id"])
    assert result is not None
    return result


@router.get("/schedule-imports", response_model=Page[ImportSummary])
async def list_imports(
    customer_id: UUID | None = None, import_status: str | None = Query(None, alias="status"),
    page: Paging = Depends(), ctx: Ctx = Depends(read),
) -> Page[ImportSummary]:
    conditions: list[str] = []
    params: dict[str, Any] = {}
    for text_, key, value in (("s.customer_id = :c", "c", customer_id),
                              ("s.status = :st", "st",
                               import_status.upper() if import_status else None)):
        if value is not None:
            conditions.append(text_)
            params[key] = value
    where = (" WHERE " + " AND ".join(conditions)) if conditions else ""
    heads = await rows(ctx, f"{IMPORT_SQL}{where} ORDER BY s.created_at DESC LIMIT :limit "
                       "OFFSET :offset", {**params, "limit": page.limit, "offset": page.offset})
    items = []
    for head in heads:
        stored = head["parsed_rows"]
        data = json.loads(stored) if isinstance(stored, str) else stored
        out = summarize(head["import_id"], head, [PreviewRow(**r) for r in data])
        items.append(ImportSummary(**out.model_dump(exclude={"rows"})))
    total = await count(ctx, f"SELECT count(*) FROM erp.schedule_import s{where}", params)
    return Page(items=items, total=total, limit=page.limit, offset=page.offset)


@router.get("/schedule-imports/{import_id}", response_model=ImportOut)
async def get_import(import_id: UUID, ctx: Ctx = Depends(read)) -> ImportOut:
    found = await load(ctx, import_id)
    if found is None:
        raise not_found("Schedule import")
    return found


async def lock_preview(ctx: Ctx, import_id: UUID) -> dict[str, Any]:
    head = await one(ctx, "SELECT import_id, customer_id, status, parsed_rows FROM "
                     "erp.schedule_import WHERE import_id = :id FOR UPDATE", {"id": import_id})
    if head is None:
        raise not_found("Schedule import")
    if head["status"] != "PREVIEW":
        raise HTTPException(status.HTTP_409_CONFLICT,
                            f"The import is {str(head['status']).lower()}")
    return head


@router.post("/schedule-imports/{import_id}/apply", response_model=ImportOut)
async def apply_import(import_id: UUID, body: ApplyIn, ctx: Ctx = Depends(manage)) -> ImportOut:
    """Create or update the delivery dates of the rows that are valid now."""
    head = await lock_preview(ctx, import_id)
    stored = head["parsed_rows"]
    data = json.loads(stored) if isinstance(stored, str) else stored
    raw_rows = [{"row_no": r["row_no"], "customer_po_no": r["customer_po_no"],
                 "customer_part_no": r["customer_part_no"],
                 "due_date": parse_date(r["due_date"]), "quantity": parse_quantity(r["quantity"])
                 if r["quantity"] is not None else None} for r in data]
    fresh = await evaluate(ctx, head["customer_id"], raw_rows)
    errors = [r for r in fresh if r.action == "ERROR"]
    if errors and not body.skip_errors:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"{len(errors)} row(s) cannot be applied; fix them or apply the valid rows only")
    applied = 0
    for r in fresh:
        if r.action == "ERROR" or r.so_line_id is None or r.due_date is None:
            continue
        await rows(
            ctx, "INSERT INTO erp.so_schedule(tenant_id, so_line_id, import_id, due_date, "
            "scheduled_qty) VALUES (erp.current_tenant_id(), :l, :i, :d, :q) ON CONFLICT "
            "(so_line_id, due_date) DO UPDATE SET scheduled_qty = EXCLUDED.scheduled_qty, "
            "import_id = EXCLUDED.import_id RETURNING so_schedule_id",
            {"l": r.so_line_id, "i": import_id, "d": r.due_date, "q": r.quantity})
        applied += 1
    await rows(
        ctx, "UPDATE erp.schedule_import SET status = 'APPLIED', applied_at = now(), applied_by = "
        ":me, parsed_rows = CAST(:rows AS jsonb) WHERE import_id = :id RETURNING import_id",
        {"me": ctx.principal.profile.app_user_id, "id": import_id,
         "rows": json.dumps([r.model_dump(mode="json") for r in fresh])})
    await audit(ctx, "schedule_import.apply", "schedule_import", import_id,
                {"applied": applied, "skipped": len(errors)})
    result = await load(ctx, import_id)
    assert result is not None
    return result


@router.post("/schedule-imports/{import_id}/discard", response_model=ImportOut)
async def discard_import(import_id: UUID, ctx: Ctx = Depends(manage)) -> ImportOut:
    await lock_preview(ctx, import_id)
    await rows(ctx, "UPDATE erp.schedule_import SET status = 'DISCARDED' WHERE import_id = :id "
               "RETURNING import_id", {"id": import_id})
    await audit(ctx, "schedule_import.discard", "schedule_import", import_id, {})
    result = await load(ctx, import_id)
    assert result is not None
    return result
